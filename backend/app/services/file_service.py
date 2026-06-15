import os
from typing import Optional


def extract_text_from_file(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext in (".pptx", ".ppt"):
        return _extract_pptx(file_path)
    elif ext in (".odp",):
        return _extract_odp_as_fallback(file_path)
    elif ext in (".txt", ".md"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    elif ext in (".png", ".jpg", ".jpeg", ".webp"):
        return f"[Image file: {filename}]"
    else:
        return ""


def _extract_pdf(file_path: str) -> str:
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(file_path)
        pages = []
        for i, page in enumerate(doc):
            text = page.get_text("text")
            if text.strip():
                pages.append(f"--- Page {i+1} ---\n{text}")
        doc.close()
        return "\n\n".join(pages)
    except Exception as e:
        return f"[PDF extraction error: {e}]"


def _extract_pptx(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        slides = []
        for i, slide in enumerate(prs.slides):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                slides.append(f"--- Slide {i+1} ---\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except Exception as e:
        return f"[PPTX extraction error: {e}]"


def _extract_odp_as_fallback(file_path: str) -> str:
    try:
        import zipfile
        from xml.etree import ElementTree as ET
        texts = []
        with zipfile.ZipFile(file_path, "r") as z:
            with z.open("content.xml") as f:
                tree = ET.parse(f)
                root = tree.getroot()
                for elem in root.iter():
                    if elem.text and elem.text.strip():
                        texts.append(elem.text.strip())
        return "\n".join(texts)
    except Exception as e:
        return f"[ODP extraction error: {e}]"
